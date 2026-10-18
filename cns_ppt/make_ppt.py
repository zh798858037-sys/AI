#!/usr/bin/env python3
"""CNS week journal club: 12 papers, one slide each + title/TOC."""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
FIG = ROOT / "figs"
OUT = ROOT / "CNS_20261015_生物医学文献汇报.pptx"

# 16:9
W, H = Inches(13.333), Inches(7.5)
FONT = "WenQuanYi Micro Hei"

NAVY = RGBColor(0x14, 0x2B, 0x45)
NAVY2 = RGBColor(0x1E, 0x4A, 0x6E)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF6, 0xF4, 0xF0)
LINE = RGBColor(0xD7, 0xDE, 0xE6)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)
CELL = RGBColor(0x00, 0x8B, 0xC7)
NATURE = RGBColor(0xC4, 0x1E, 0x3A)
SCIENCE = RGBColor(0x00, 0x3D, 0x7A)
GREEN = RGBColor(0x1F, 0x7A, 0x4D)

JOURNAL_COLOR = {"Cell": CELL, "Nature": NATURE, "Science": SCIENCE}


def set_run_font(run, size, bold=False, color=INK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)
    cs = rPr.find(qn("a:cs"))
    if cs is None:
        cs = rPr.makeelement(qn("a:cs"), {})
        rPr.append(cs)
    cs.set("typeface", name)


def add_textbox(slide, l, t, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    run = p.add_run()
    run.text = text
    set_run_font(run, size, bold=bold, color=color)
    return box


def add_rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def add_round(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    # tighter corners
    try:
        sh.adjustments[0] = 0.12
    except Exception:
        pass
    return sh


def add_picture_fit(slide, path, l, t, w, h):
    """Place image inside box, keep aspect ratio, centered."""
    im = Image.open(path)
    iw, ih = im.size
    box_aspect = w / h
    img_aspect = iw / ih
    if img_aspect > box_aspect:
        nw, nh = w, int(w / img_aspect)
    else:
        nh, nw = h, int(h * img_aspect)
    left = l + (w - nw) // 2
    top = t + (h - nh) // 2
    return slide.shapes.add_picture(str(path), left, top, nw, nh)


def make_dendrite_schematic(path: Path):
    """Main-concept schematic for the Science dendritic computation paper."""
    Wpx, Hpx = 1600, 1100
    img = Image.new("RGB", (Wpx, Hpx), (255, 255, 255))
    d = ImageDraw.Draw(img)
    font_path = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"
    f_title = ImageFont.truetype(font_path, 42)
    f_h = ImageFont.truetype(font_path, 30)
    f_b = ImageFont.truetype(font_path, 24)
    f_s = ImageFont.truetype(font_path, 20)

    d.rounded_rectangle([40, 30, 1560, 100], 12, fill=(0, 61, 122))
    d.text((60, 48), "同一 L5 ET 神经元内：树突编码规则，胞体编码感觉", font=f_title, fill="white")

    # neuron area
    d.rounded_rectangle([60, 130, 980, 1040], 18, fill=(246, 249, 252), outline=(215, 222, 230), width=2)

    # apical tuft
    d.ellipse([250, 170, 790, 310], fill=(232, 244, 255), outline=(0, 139, 199), width=4)
    d.text((360, 210), "顶端树突  Apical dendrite", font=f_h, fill=(0, 90, 150))
    d.text((390, 250), "优先编码：推断出的抽象规则", font=f_b, fill=(0, 90, 150))

    # trunk
    d.rectangle([490, 305, 550, 620], fill=(176, 214, 236))
    d.polygon([(520, 305), (430, 430), (460, 430), (520, 360), (580, 430), (610, 430)], fill=(176, 214, 236))
    d.polygon([(520, 500), (400, 600), (430, 600), (520, 530), (610, 600), (640, 600)], fill=(176, 214, 236))

    # soma
    d.ellipse([390, 600, 650, 840], fill=(255, 236, 220), outline=(196, 92, 38), width=5)
    d.text((445, 680), "胞体  Soma", font=f_h, fill=(154, 62, 16))
    d.text((415, 725), "优先编码：感觉刺激", font=f_b, fill=(154, 62, 16))
    # axon
    d.rectangle([505, 838, 535, 980], fill=(196, 92, 38))
    d.polygon([(490, 975), (550, 975), (520, 1020)], fill=(196, 92, 38))

    # choice badges
    d.rounded_rectangle([70, 360, 280, 450], 12, fill=(31, 122, 77))
    d.text((92, 375), "Choice 信号", font=f_b, fill="white")
    d.text((100, 408), "树突也有", font=f_s, fill="white")
    d.line([(280, 405), (490, 280)], fill=(31, 122, 77), width=4)

    d.rounded_rectangle([700, 700, 940, 790], 12, fill=(31, 122, 77))
    d.text((722, 715), "Choice 信号", font=f_b, fill="white")
    d.text((730, 748), "胞体也有", font=f_s, fill="white")
    d.line([(700, 745), (640, 730)], fill=(31, 122, 77), width=4)

    # right panels
    d.rounded_rectangle([1020, 130, 1540, 470], 16, fill=(232, 244, 255), outline=(0, 139, 199), width=2)
    d.text((1045, 155), "任务设计", font=f_h, fill=(0, 61, 122))
    d.text((1045, 210), "• 规则切换的听觉灵活分类", font=f_b, fill=(31, 41, 55))
    d.text((1045, 255), "• 边界无预警变化，小鼠需从", font=f_b, fill=(31, 41, 55))
    d.text((1045, 295), "  奖励反馈中推断隐藏规则", font=f_b, fill=(31, 41, 55))
    d.text((1045, 345), "• 双光子成像：同步记录", font=f_b, fill=(31, 41, 55))
    d.text((1045, 385), "  同一神经元树突与胞体", font=f_b, fill=(31, 41, 55))
    d.text((1045, 425), "• 光遗传抑制树突检验因果", font=f_b, fill=(31, 41, 55))

    d.rounded_rectangle([1020, 500, 1540, 1040], 16, fill=(255, 243, 232), outline=(196, 92, 38), width=2)
    d.text((1045, 525), "关键因果证据", font=f_h, fill=(154, 62, 16))
    d.text((1045, 585), "光遗传抑制顶端树突", font=f_b, fill=(31, 41, 55))
    d.text((1045, 630), "→ 胞体规则依赖的选择编码", font=f_b, fill=(31, 41, 55))
    d.text((1045, 670), "  明显减弱", font=f_b, fill=(31, 41, 55))
    d.text((1045, 730), "结论：树突–胞体跨区室整合", font=f_b, fill=(154, 62, 16))
    d.text((1045, 775), "是灵活决策的细胞计算基础。", font=f_b, fill=(154, 62, 16))
    d.text((1045, 840), "单个皮层神经元内部即可完成", font=f_s, fill=(75, 85, 99))
    d.text((1045, 875), "“规则 × 感觉 → 选择”的认知运算。", font=f_s, fill=(75, 85, 99))
    d.text((1045, 940), "Zhang / Xu  Ning-long  Science 2026", font=f_s, fill=(107, 114, 128))

    img.save(path, "PNG")


def add_footer(slide, n, total=12):
    add_rect(slide, 0, Inches(7.22), W, Inches(0.28), NAVY)
    add_textbox(
        slide,
        Inches(0.35),
        Inches(7.22),
        Inches(9),
        Inches(0.28),
        "CNS 三大主刊 2026.10.15 ｜生物医学文献汇报（一文一页）",
        size=10,
        color=WHITE,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_textbox(
        slide,
        Inches(10.6),
        Inches(7.22),
        Inches(2.4),
        Inches(0.28),
        f"{n} / {total}",
        size=10,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def add_bullets(slide, l, t, w, h, items, size=13, color=INK, spacing=6):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        p.level = 0
        run = p.add_run()
        run.text = "•  " + item
        set_run_font(run, size, color=color)
    return box


def paper_slide(prs, meta, idx, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, W, H, CREAM)
    add_rect(slide, 0, 0, W, Inches(1.08), NAVY)

    badge_c = JOURNAL_COLOR[meta["journal"]]
    add_round(slide, Inches(0.32), Inches(0.22), Inches(1.35), Inches(0.34), badge_c)
    add_textbox(
        slide,
        Inches(0.32),
        Inches(0.22),
        Inches(1.35),
        Inches(0.34),
        meta["journal"],
        size=12,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_textbox(
        slide,
        Inches(1.8),
        Inches(0.14),
        Inches(11.2),
        Inches(0.42),
        meta["title_cn"],
        size=18,
        bold=True,
        color=WHITE,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_textbox(
        slide,
        Inches(1.8),
        Inches(0.56),
        Inches(11.2),
        Inches(0.42),
        meta["title_en"],
        size=11,
        color=RGBColor(0xD0, 0xDC, 0xE8),
        anchor=MSO_ANCHOR.TOP,
    )

    # left column
    lx, ly, lw = Inches(0.32), Inches(1.22), Inches(6.15)

    add_round(slide, lx, ly, lw, Inches(1.18), WHITE)
    add_textbox(slide, lx + Inches(0.16), ly + Inches(0.08), lw - Inches(0.28), Inches(0.28), "解决什么问题", size=12, bold=True, color=badge_c)
    add_textbox(slide, lx + Inches(0.16), ly + Inches(0.36), lw - Inches(0.28), Inches(0.74), meta["problem"], size=13, color=INK)

    add_round(slide, lx, Inches(2.52), lw, Inches(1.55), WHITE)
    add_textbox(slide, lx + Inches(0.16), Inches(2.58), lw - Inches(0.28), Inches(0.28), "做了什么", size=12, bold=True, color=badge_c)
    add_textbox(slide, lx + Inches(0.16), Inches(2.86), lw - Inches(0.28), Inches(1.12), meta["did"], size=13, color=INK)

    add_round(slide, lx, Inches(4.18), lw, Inches(1.95), WHITE)
    add_textbox(slide, lx + Inches(0.16), Inches(4.24), lw - Inches(0.28), Inches(0.28), "关键结论", size=12, bold=True, color=badge_c)
    add_bullets(slide, lx + Inches(0.12), Inches(4.52), lw - Inches(0.24), Inches(1.52), meta["findings"], size=13)

    add_textbox(
        slide,
        lx,
        Inches(6.18),
        lw,
        Inches(0.92),
        meta["authors"] + "\n" + meta["cite"],
        size=10,
        color=MUTED,
    )

    # right figure panel
    fx, fy, fw, fh = Inches(6.62), Inches(1.22), Inches(6.38), Inches(5.85)
    add_round(slide, fx, fy, fw, fh, WHITE)
    add_textbox(slide, fx + Inches(0.16), fy + Inches(0.06), fw - Inches(0.3), Inches(0.28), meta["fig_label"], size=11, bold=True, color=MUTED)
    add_picture_fit(
        slide,
        meta["fig"],
        fx + Inches(0.12),
        fy + Inches(0.34),
        fw - Inches(0.24),
        fh - Inches(0.46),
    )

    add_footer(slide, idx, total)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, W, H, NAVY)
    add_rect(slide, 0, Inches(5.9), W, Inches(1.6), NAVY2)
    add_textbox(slide, Inches(0.7), Inches(1.5), Inches(12), Inches(0.4), "CNS 三大主刊 ｜ 2026 年 10 月 15 日当周", size=16, color=RGBColor(0xC4, 0xD4, 0xE6))
    add_textbox(slide, Inches(0.7), Inches(2.05), Inches(12), Inches(1.3), "生物医学文献汇报", size=40, bold=True, color=WHITE)
    add_textbox(
        slide,
        Inches(0.7),
        Inches(3.4),
        Inches(12),
        Inches(1.2),
        "12 篇研究论文  ·  一文一页  ·  问题 / 工作 / 结论 + 主图",
        size=18,
        color=RGBColor(0xE8, 0xEE, 0xF4),
    )
    add_textbox(
        slide,
        Inches(0.7),
        Inches(6.15),
        Inches(12),
        Inches(1.1),
        "Cell  Vol.189 Issue 21    Nature  Vol.658 Issue 8136    Science  Vol.394 Issue 6821\n"
        "范围：研究论文（Article）。主图优先采用 Graphical Abstract 或 Figure 1。",
        size=13,
        color=RGBColor(0xC4, 0xD4, 0xE6),
    )


def toc_slide(prs, papers):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, W, H, CREAM)
    add_rect(slide, 0, 0, W, Inches(0.85), NAVY)
    add_textbox(slide, Inches(0.45), Inches(0.18), Inches(12), Inches(0.5), "目录  ·  12 篇一文一页", size=22, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

    groups = [
        ("你点名的 3 篇", papers[:3], CELL),
        ("较高相关 5 篇（递送 / NDD 窗口 / 皮层发育）", papers[3:8], NATURE),
        ("中等相关 4 篇（胶质 / 神经保护 / AD / 树突计算）", papers[8:], SCIENCE),
    ]
    y = Inches(1.05)
    n = 1
    for gtitle, items, col in groups:
        add_textbox(slide, Inches(0.5), y, Inches(12), Inches(0.32), gtitle, size=14, bold=True, color=col)
        y += Inches(0.34)
        for p in items:
            add_round(slide, Inches(0.5), y, Inches(0.42), Inches(0.32), col)
            add_textbox(slide, Inches(0.5), y, Inches(0.42), Inches(0.32), f"{n:02d}", size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            add_textbox(
                slide,
                Inches(1.05),
                y,
                Inches(11.7),
                Inches(0.32),
                f"[{p['journal']}]  {p['title_cn']}",
                size=13,
                color=INK,
                anchor=MSO_ANCHOR.MIDDLE,
            )
            y += Inches(0.36)
            n += 1
        y += Inches(0.08)
    add_footer(slide, 0, 12)


PAPERS = [
    dict(
        journal="Cell",
        title_cn="用蛋白降解驱动比例型转录激活，把“蛋白变少”变成可筛选信号",
        title_en="Ratiometric transcriptional activation by protein degradation  ·  Gray / Banik",
        problem="靶向蛋白降解药（PROTAC、分子胶）很难在活细胞里把“蛋白略降”放大成稳定、可多路过的筛选读出；常规检测是蛋白越少信号越弱。",
        did="设计 RTA 遗传回路：把目标蛋白（POI）与转录因子抑制剂融合。POI 被降解后抑制剂减少，转录因子被释放。灵敏度由 POI-抑制剂 / 转录因子的比例决定，而不是蛋白绝对量。",
        findings=[
            "蛋白下降可被放大为 100–1000× 转录输出；不依赖具体降解机制。",
            "可在单细胞做三路正交回路，也可在细胞池用荧光、存活或 barcode 测序读出。",
            "用于癌基因靶点筛选，可快速发现新的分子胶降解剂。",
        ],
        authors="Melissa A. Gray, Steven M. Banik 等  ·  Stanford",
        cite="Cell 189, 6666–6683 (2026)  ·  DOI: 10.1016/j.cell.2026.08.009",
        fig=FIG / "rta_ga.jpg",
        fig_label="主图  Graphical abstract（来源：Cell）",
    ),
    dict(
        journal="Cell",
        title_cn="雌激素重编程 STING，促使 NET“功能性抗原屏蔽”，造成 ADC 的女性偏倚耐药",
        title_en="STING reprogramming by estrogen … sex-biased resistance to ADC therapy  ·  Yan / Hao",
        problem="抗体偶联药物（ADC）疗效是否存在性别差异、机制是什么，此前不清楚；抗原表达相同却出现女性耐药更难解释。",
        did="用 55 例 CLDN18.2+ 胰腺癌人源化 PDX 做“类 II 期临床前试验”，比较雌雄宿主对 CLDN18.2-ADC 的响应，并解析雌激素–GPER1–TRAF6–STING 通路。",
        findings=[
            "抗原水平相当，但雌性宿主对 CLDN18.2-ADC 显著耐药。",
            "雌激素经 GPER1 上调 TRAF6，把 STING 从 IRF3/IFN 转到 NF-κB/CXCL1，招募中性粒细胞形成 NET。",
            "NET 构成带电胞外网，限制抗体接触抗原（功能性抗原屏蔽）。Sivelestat 破坏 NET 后雌性疗效恢复；临床队列中 NET 负荷也与多种抗体疗法疗效差相关。",
        ],
        authors="闫璟瑞、周天兴、石尚亨、王滨、郝继辉 等  ·  天津医科大学肿瘤医院",
        cite="Cell 189, 6752–6770 (2026)  ·  DOI: 10.1016/j.cell.2026.09.031",
        fig=FIG / "sting_ga.jpg",
        fig_label="主图  Graphical abstract（来源：Cell）",
    ),
    dict(
        journal="Cell",
        title_cn="B 细胞“征用”古老 LINE-1，借助 HUSH 把 AID 精确带到 VH 基因",
        title_en="Co-option of retrotransposons promotes antibody diversification  ·  Lauring / Basu",
        problem="体细胞超突变（SHM）依赖 AID，但 AID 如何特异打到重排后的 VH(D)JH，而不是基因组其他地方，一直是抗体多样性的核心谜题。",
        did="发现多数 VH 基因上游存在古老、截短、已无转座能力的 LINE-1；证明重排 VH 启动子产生编码上游 L1 的长反义 RNA，招募 HUSH 复合体与 AID。",
        findings=[
            "提出 L1-driven SHM：上游 L1 在顺式招募 HUSH/AID。",
            "HUSH 条件敲除或去掉 VH 上游 L1 的小鼠，体内 SHM 下降。",
            "把有转录活性的 L1 插到非 Ig 位点会引发脱靶 SHM。B 细胞把抗转座子沉默机制改造成抗体多样化工具，连接了固有免疫与适应性免疫。",
        ],
        authors="Max C. Lauring, Jiguang Wang, Uttiya Basu 等  ·  Columbia / HKUST",
        cite="Cell 189, 6684–6704 (2026)  ·  DOI: 10.1016/j.cell.2026.07.021",
        fig=FIG / "l1_ga.jpg",
        fig_label="主图  Graphical abstract（来源：Cell）",
    ),
    dict(
        journal="Nature",
        title_cn="AI 设计合成蛋白组装体，做成比病毒/LNP 更强的 RNA 递送载具",
        title_en="Creating bottom-up RNA transfer vehicles from synthetic protein assemblies  ·  Schuhmacher / Wurst",
        problem="天然病毒衣壳受进化约束，尺寸/形状趋同；临床 LNP 递送效率与靶向性仍有限。能否完全抛开天然衣壳，用生成式蛋白设计从头做 RNA 运输系统？",
        did="把 RFdiffusion 设计的合成蛋白组装体，与膜结合、出芽、RNA 结合结构域融合，筛 >100 种 bottom-up 合成转运载具（STV）。最优为平面 C8 对称的 STV-C8。",
        findings=[
            "STV-C8 的 RNA 递送效率比常用病毒样颗粒和临床 LNP 高数个数量级。",
            "可用计算设计的肽结合体编程嗜性；可运载报告 RNA、基因编辑器、抗病毒 RNA、转录因子。",
            "小鼠近单细胞分辨率 biodistribution 显示安全性良好；在患者来源细胞和猪中递送 Cas9，删除 dystrophin 第 51 外显子（DMD 策略）。",
        ],
        authors="Maren Kirstin Schuhmacher, Wolfgang Wurst 等  ·  Helmholtz / TUM",
        cite="Nature 658, 805–814 (2026)  ·  DOI: 10.1038/s41586-026-10952-3",
        fig=FIG / "stv_fig1.png",
        fig_label="主图  Figure 1（来源：Nature）",
    ),
    dict(
        journal="Nature",
        title_cn="Binding-to-release：结合口袋里就裂解，靶向递药不再依赖内吞",
        title_en="A binding-to-release strategy for targeted anticancer drug delivery  ·  Wen / Liu",
        problem="ADC/SMDC 大多走 internalization-to-release（ITR）。约 90% 膜蛋白内吞差，结合后药物很快解离被清除，可用靶点被严重限制。",
        did="提出 binding-to-release（BTR）：配体结合后，连接子被结合口袋里邻近亲核残基直接切断。开发磷(V)–酚交换化学 PhoPEx（受 SuFEx 启发），实现循环稳定、邻近触发释放。",
        findings=[
            "FAP 的 Tyr745 可触发裂解；无需内吞即可在瘤灶释放可通透细胞的 payload。",
            "FAP-BTR-SMDC 瘤内 MMAE 暴露（AUC0–120 h）是内吞依赖 SMDC 的 5.9 倍，瘤/血、瘤/肝比显著优于 ITR-SMDC 和 ITR-ADC。",
            "多种瘤模型接近完全消退；并扩展到 PD-L1 及 mRNA display 得到的 FAP 肽，扩大靶点范围。",
        ],
        authors="文梓豪、徐梦欣、刘志博 等  ·  北京大学",
        cite="Nature 658, 785–794 (2026)  ·  DOI: 10.1038/s41586-026-10971-0",
        fig=FIG / "btr_fig1.png",
        fig_label="主图  Figure 1（来源：Nature）",
    ),
    dict(
        journal="Nature",
        title_cn="成体恢复海马 PV 中间神经元可塑性，可逆转神经发育障碍的认知与癫痫表型",
        title_en="Procognitive restoration of PV neuron plasticity in neurodevelopmental disorders  ·  Shih / Sahay",
        problem="NDD（自闭、癫痫等）常被认为是发育窗口关闭后不可逆。PV 中间神经元早期受损导致网络过兴奋，成年还能不能把环路和认知拉回来？",
        did="用输入特异的转译组筛选，鉴定成体 CA3/CA2 中“经验依赖 PV 可塑性基因”（XPG）。大量 XPG 在自闭/癫痫/双相/精分中呈单倍不足。在成年 Cntnap2−/− 小鼠的 PV 神经元中上调候选基因 Meis2。",
        findings=[
            "Meis2 足以恢复经验依赖的 PV 可塑性、兴奋–抑制平衡。",
            "改善空间/社会记忆、神经元集群与尖波涟漪，并抑制癫痫。",
            "说明部分发育缺陷在成年仍可逆；关键是对准正确细胞类型的可塑性程序，而不只是“补一个全局信号”。",
        ],
        authors="Yu-Tzu Shih, Amar Sahay 等  ·  MGH / Harvard",
        cite="Nature 658, 725 附近 (2026)  ·  DOI: 10.1038/s41586-026-10907-8",
        fig=FIG / "pv_fig1.png",
        fig_label="主图  Figure 1（来源：Nature）",
    ),
    dict(
        journal="Nature",
        title_cn="皮层类器官里放射状胶质细胞“会报时、不会报分”：谱系时序与体内脱钩",
        title_en="Temporal uncoupling of radial glia lineage progression in cortical organoids  ·  Stouffer / Hippenmeyer",
        problem="自组织皮层类器官常被当作人脑发育模型，但放射状胶质祖细胞（RGP）是否仍按体内那套严格时序推进谱系，缺乏单克隆定量证据。",
        did="把 MADM 谱系示踪做到小鼠 ESC 来源皮层类器官，与体内 RGP 定量框架直接对比：增殖潜能、克隆神经元多样性、SCPN vs CPN 命运。",
        findings=[
            "类器官 RGP 增殖潜能高度可塑，不再是体内那种刻板的时序谱系推进。",
            "约 1/3 克隆出现命运限制（体内仅约 3%），克隆内皮层投射神经元多样性下降。",
            "单细胞转录轨迹仍均一，说明缺的是非细胞自主龛位信号。筛药/建模时要分清“药物无效”还是“模型缺时序”。",
        ],
        authors="Melissa Stouffer, Simon Hippenmeyer 等  ·  ISTA",
        cite="Nature 658, 716–724 (2026)  ·  DOI: 10.1038/s41586-026-10916-7",
        fig=FIG / "rgp_fig1.png",
        fig_label="主图  Figure 1（来源：Nature）",
    ),
    dict(
        journal="Cell",
        title_cn="早期人皮层代谢图谱：糖酵解/磷酸戊糖途径开关决定放射状胶质命运",
        title_en="Metabolic atlas of early human cortex … PPP control of cell fate  ·  Mil / Bhaduri",
        problem="人皮层发育需要快速扩增祖细胞并多样化，但人源组织中代谢程序几乎未被系统画过，代谢如何卡细胞命运也不清。",
        did="用原代人皮层 + 皮层类器官做多模态代谢图谱。发现晚期神经发生时糖酵解和磷酸戊糖途径（PPP）意外升高，再用降糖、药理/遗传抑制 PPP、补核糖做干预。",
        findings=[
            "降糖或抑制 PPP 酶 → 外侧放射状胶质（oRG）和抑制性神经元增多。",
            "核糖足以挽救 RG 基因表达、类器官细胞组成，并恢复 ATP 与亚牛磺酸。",
            "PPP 是调控 RG 命运的关键节点；为发育代谢与疾病模型提供资源。",
        ],
        authors="Jessenya Mil, Heather Christofk, Aparna Bhaduri 等  ·  UCLA",
        cite="Cell 189, 6810–6828 (2026)  ·  DOI: 10.1016/j.cell.2026.07.023",
        fig=FIG / "metab_ga.jpg",
        fig_label="主图  Graphical abstract（来源：Cell）",
    ),
    dict(
        journal="Nature",
        title_cn="发育中人皮层四种胶质细胞的 3D 表观基因组，定位精分风险与人加速区",
        title_en="3D epigenome of glial cell types in developing human cortex  ·  Jones / Shen",
        problem="非编码变异如何通过胶质细胞特异调控影响神经精神疾病和人特有皮层扩张，缺乏中孕期原代胶质的 3D 表观组图谱。",
        did="FACS 分选中孕期人新皮层 vRG、oRG、OPC、小胶质细胞，整合 RNA-seq、ATAC、WGBS、PLAC-seq（H3K4me3），鉴定细胞类型特异 cCRE，并用转基因小鼠胚胎验证。",
        findings=[
            "机器学习优先排序 112 个落在胶质 cCRE 的精分风险变异；rs4449074 可在体内破坏预测的 vRG 增强子。",
            "oRG 的 cCRE 富集人加速区（HAR），部分 HAR 与黑猩猩直向同源序列活性不同，并与神经发育基因互作。",
            "为胶质来源的疾病与人特异皮层进化提供调控地图。",
        ],
        authors="Ian R. Jones, Li Wang, Yin Shen, Arnold Kriegstein 等  ·  UCSF",
        cite="Nature 658, 795–804 (2026)  ·  DOI: 10.1038/s41586-026-10987-6",
        fig=FIG / "glia_fig1.png",
        fig_label="主图  Figure 1（来源：Nature）",
    ),
    dict(
        journal="Nature",
        title_cn="补体 H 因子在神经元细胞内抗铁死亡，保护炎症中的中枢神经元",
        title_en="Intracellular complement factor H protects neurons during CNS inflammation  ·  Mayer / Friese",
        problem="多发性硬化的残疾主要由神经退变驱动，但神经元亚型为何有的耐炎、有的易损，神经元内在机制不清。CFH 一直被当成胞外补体调节蛋白。",
        did="对 MS 供者视网膜神经节细胞做 snRNA-seq，发现 CFH 与内在耐受力强相关。再在 EAE 小鼠中做神经元特异 Cfh 缺失/过表达，并解析亚细胞定位与结构域。",
        findings=[
            "炎症/氧化应激诱导神经元表达 CFH；它定位内质网，限制 ROS 和脂质过氧化（铁死亡相关）。",
            "保护依赖 C 端 SCR20，不依赖分泌，也不依赖 C3。",
            "神经元缺 CFH 使 EAE 更重、视觉更差、神经元丢失更多；胞内 CFH/SCR20 可改善恢复。提出非经典的胞内神经保护功能。",
        ],
        authors="Christina Mayer, Marcel S. Woo, Manuel A. Friese 等  ·  Hamburg INIMS",
        cite="Nature 658, 765–773 (2026)  ·  DOI: 10.1038/s41586-026-10981-y",
        fig=FIG / "cfh_fig1.png",
        fig_label="主图  Figure 1（来源：Nature）",
    ),
    dict(
        journal="Nature",
        title_cn="兴奋性神经元异常表达 ERBB4 是阿尔茨海默病早期病理的驱动因素",
        title_en="Aberrant excitatory neuronal ERBB4 promotes Alzheimer’s disease pathology  ·  Lee / Chung",
        problem="AD 早期突触丢失常被归为小胶质过度吞噬。但胶质吞噬异常究竟是因还是果？神经炎症是否足以解释早期突触失衡？",
        did="在 AD 小鼠中定量胶质对兴奋/抑制突触的吞噬，并用 snRNA-seq 抓住最早出现的“早响应兴奋性神经元（EREN）”，其标志是异位表达 Erbb4。再做兴奋性神经元特异敲除/过表达。",
        findings=[
            "星形胶质/小胶质增加清除兴奋性突触、减少清除抑制性突触；单有神经炎症不足以解释早期突触丢失。",
            "兴奋性神经元敲除 Erbb4：网络活动、突触平衡、胶质增生、淀粉样斑块和认知都改善。",
            "野生型里过表达 Erbb4 即可重现核心 AD 样表型，且不需要斑块；下游依赖 mTOR。人转录组中介分析支持 ERBB4 连接淀粉样→tau→认知下降。",
        ],
        authors="Se Young Lee, Won-Suk Chung 等  ·  KAIST",
        cite="Nature 658, 725 附近 (2026)  ·  DOI: 10.1038/s41586-026-10964-z",
        fig=FIG / "erbb4_fig1.png",
        fig_label="主图  Figure 1（来源：Nature）",
    ),
    dict(
        journal="Science",
        title_cn="灵活决策的细胞计算：顶端树突编码抽象规则，胞体编码感觉",
        title_en="Dendritic computation for abstract rule-based flexible decision-making  ·  Zhang / Xu",
        problem="智能行为要把感觉输入和不断变化的抽象规则整合。树突计算已知参与感觉的自上而下调制，但有没有直接证据表明它参与灵活认知决策？",
        did="训练小鼠做规则切换的听觉灵活分类，用双光子成像同步记录听皮层 L5 端脑外投射（ET）神经元的顶端树突与胞体，再用光遗传抑制树突做因果检验。",
        findings=[
            "同一神经元内功能分区：树突优先编码推断规则，胞体优先表征感觉；选择信号两边都有。",
            "抑制树突会损害胞体上规则依赖的选择编码，证明跨区室整合是任务必需的。",
            "单个皮层神经元内部就能完成“规则 × 感觉 → 选择”，把树突双区室模型从感觉运动拓展到抽象认知。",
        ],
        authors="张园、徐宁龙 等  ·  中科院脑智卓越中心",
        cite="Science 394, 287 (2026)  ·  DOI: 10.1126/science.adz5323",
        fig=FIG / "dend_scheme.png",
        fig_label="主图  按原文核心结论重绘的机制示意图（Science 原图因版权墙未能下载）",
    ),
]


def main():
    dend_path = FIG / "dend_scheme.png"
    make_dendrite_schematic(dend_path)

    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    title_slide(prs)
    toc_slide(prs, PAPERS)
    for i, meta in enumerate(PAPERS, 1):
        if not Path(meta["fig"]).exists():
            raise FileNotFoundError(meta["fig"])
        paper_slide(prs, meta, i, len(PAPERS))
    prs.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
